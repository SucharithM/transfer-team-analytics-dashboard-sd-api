import argparse
import json
import os
import sqlite3
from pathlib import Path

from openpyxl import load_workbook
import pandas as pd

from workflow_dashboard import __main__ as cli
from workflow_dashboard import generation
from workflow_dashboard.config import DashboardConfig
from workflow_dashboard.history import build_metric_definition, store_snapshot

from .test_analytics_cohorts import build_sample_analytics
from .test_analytics_cohorts import sample_cases


def _run_args(
    tmp_path: Path, *, output_dir: Path | None = None
) -> argparse.Namespace:
    return argparse.Namespace(
        source="fixture",
        output_dir=output_dir or tmp_path / "outputs",
        fixture_path=Path("data/sample_api_response.json"),
        history_db=tmp_path / "history.sqlite",
    )


def _output_files(output_dir: Path) -> list[Path]:
    if not output_dir.is_dir():
        return []
    return [path for path in output_dir.rglob("*") if path.is_file()]


def _snapshot_count(db_path: Path) -> int:
    with sqlite3.connect(db_path) as connection:
        return connection.execute(
            "SELECT COUNT(*) FROM snapshot_metrics"
        ).fetchone()[0]


def _sample_definition():
    return build_metric_definition(
        DashboardConfig(),
        sample_cases(),
        source="fixture",
    )


def test_output_directory_failure_writes_no_history(tmp_path, capsys) -> None:
    output_file = tmp_path / "output-is-a-file"
    output_file.write_text("not a directory", encoding="utf-8")
    args = _run_args(tmp_path, output_dir=output_file)

    assert cli.run_command(args) == 2
    output = capsys.readouterr().out
    assert "Check file access" in output
    assert str(output_file) not in output

    assert not args.history_db.exists()


def test_late_writer_failure_publishes_nothing_and_preserves_history(
    monkeypatch, tmp_path, capsys
) -> None:
    args = _run_args(tmp_path)
    store_snapshot(
        args.history_db,
        build_sample_analytics(),
        _sample_definition(),
        source="fixture",
    )

    def fail_workbook(**kwargs) -> None:
        raise RuntimeError("workbook failed")

    monkeypatch.setattr(generation, "export_workbook", fail_workbook)

    assert cli.run_command(args) == 2
    output = capsys.readouterr().out
    assert json.loads(output)["error"]
    assert "workbook failed" not in output

    assert _snapshot_count(args.history_db) == 1
    assert _output_files(args.output_dir) == []
    assert not any(
        path.name.startswith(".workflow-run-")
        for path in args.output_dir.rglob("*")
    )


def test_partial_publication_is_removed_before_history_commit(
    monkeypatch, tmp_path, capsys
) -> None:
    args = _run_args(tmp_path)
    real_replace = os.replace
    replace_calls = 0

    def fail_second_replace(source, destination) -> None:
        nonlocal replace_calls
        replace_calls += 1
        if replace_calls == 2:
            raise OSError("publish failed")
        real_replace(source, destination)

    monkeypatch.setattr(generation.os, "replace", fail_second_replace)

    assert cli.run_command(args) == 2
    assert "publish failed" not in capsys.readouterr().out

    assert replace_calls == 2
    assert not args.history_db.exists()
    assert _output_files(args.output_dir) == []


def test_history_failure_removes_published_artifacts(monkeypatch, tmp_path, capsys) -> None:
    args = _run_args(tmp_path)

    def fail_history(*args, **kwargs) -> None:
        raise sqlite3.OperationalError("history failed")

    monkeypatch.setattr(generation, "store_prepared_snapshot", fail_history)

    assert cli.run_command(args) == 2
    assert "history failed" not in capsys.readouterr().out

    assert not args.history_db.exists()
    assert _output_files(args.output_dir) == []


def test_success_publishes_complete_set_and_commits_rendered_run(
    tmp_path, capsys
) -> None:
    args = _run_args(tmp_path)

    assert cli.run_command(args) == 0
    result = json.loads(capsys.readouterr().out)

    output_paths = {
        name: Path(path)
        for name, path in result["outputs"].items()
        if name != "history_db"
    }
    assert set(output_paths) == {"csv", "html", "excel"}
    assert all(
        path.is_file() and path.stat().st_size > 0
        for path in output_paths.values()
    )
    assert _snapshot_count(args.history_db) == 1

    with sqlite3.connect(args.history_db) as connection:
        owner_count = connection.execute(
            "SELECT COUNT(*) FROM owner_workload_history WHERE run_id = ?",
            (result["run_id"],),
        ).fetchone()[0]
    assert owner_count == 6

    workbook = load_workbook(output_paths["excel"], read_only=True)
    history_rows = list(workbook["History Trends"].iter_rows(values_only=True))
    run_id_column = history_rows[0].index("run_id")
    assert result["run_id"] in {
        row[run_id_column] for row in history_rows[1:]
    }


def _mixed_workflow_records() -> list[dict[str, str]]:
    base = {
        "stepName": "Review",
        "status": "InProgress",
        "submittedBy": "Staff Member A",
        "submissionDate": "2026-07-01T12:00:00Z",
        "lastActivityDate": "2026-07-20T12:00:00Z",
    }
    return [
        {
            **base,
            "packageID": "included",
            "packageName": "Included package",
            "workflowName": "Workflow A",
        },
        {
            **base,
            "packageID": "excluded",
            "packageName": "Excluded package",
            "workflowName": "Workflow B",
        },
    ]


def _fixture_fetch_meta(record_count: int) -> dict[str, object]:
    return {
        "pages": 1,
        "attempts": 1,
        "requested_take": None,
        "reported_total": record_count,
        "retrieved_records": record_count,
        "unique_records": None,
        "duplicate_records": None,
        "pagination_validated": False,
    }


def test_population_validation_failure_writes_nothing(
    monkeypatch, tmp_path, capsys
) -> None:
    args = _run_args(tmp_path)
    config = DashboardConfig(
        output_dir=args.output_dir,
        fixture_path=args.fixture_path,
        history_db_path=args.history_db,
    )
    records = _mixed_workflow_records()
    monkeypatch.setattr(cli, "load_config", lambda **kwargs: config)
    monkeypatch.setattr(
        cli,
        "records_from_source",
        lambda source, config: (records, _fixture_fetch_meta(len(records)), {}),
    )

    assert cli.run_command(args) == 2

    result = json.loads(capsys.readouterr().out)
    assert "Check the configured workflow" in result["error"]
    assert all(record["workflowName"] not in result["error"] for record in records)
    assert not args.output_dir.exists()
    assert not args.history_db.exists()


def test_filtered_population_reaches_every_artifact_and_history(
    monkeypatch, tmp_path, capsys
) -> None:
    args = _run_args(tmp_path)
    config = DashboardConfig(
        workflow_filter="workflow a",
        output_dir=args.output_dir,
        fixture_path=args.fixture_path,
        history_db_path=args.history_db,
    )
    records = _mixed_workflow_records()
    monkeypatch.setattr(cli, "load_config", lambda **kwargs: config)
    monkeypatch.setattr(
        cli,
        "records_from_source",
        lambda source, config: (records, _fixture_fetch_meta(len(records)), {}),
    )

    assert cli.run_command(args) == 0

    result = json.loads(capsys.readouterr().out)
    assert result["fetch"]["retrieved_records"] == 2
    assert result["records"] == 1

    normalized = pd.read_csv(result["outputs"]["csv"])
    assert normalized["record_id"].tolist() == ["included"]
    assert normalized["workflow_name"].tolist() == ["Workflow A"]

    workbook = load_workbook(result["outputs"]["excel"], read_only=True)
    normalized_rows = list(
        workbook["Normalized Records"].iter_rows(values_only=True)
    )
    record_id_column = normalized_rows[0].index("package_id")
    assert [row[record_id_column] for row in normalized_rows[1:]] == ["included"]

    with sqlite3.connect(args.history_db) as connection:
        stored_count = connection.execute(
            "SELECT total_records FROM snapshot_metrics"
        ).fetchone()[0]
        definition_json = connection.execute(
            "SELECT definition_json FROM metric_definitions"
        ).fetchone()[0]
    assert stored_count == 1
    assert json.loads(definition_json)["population"] == {
        "includes_unnamed_workflow": False,
        "source": "fixture",
        "source_identity": str(args.fixture_path.resolve()),
        "workflow_filter": "workflow a",
        "workflow_names": ["Workflow A"],
    }
