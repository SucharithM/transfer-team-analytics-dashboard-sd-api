import sqlite3

import numpy as np
from openpyxl import load_workbook
import pandas as pd
import pytest

from workflow_dashboard.generation import _csv_ready
from workflow_dashboard.analytics import (
    _closed_by_month,
    _daily_closures,
    build_analytics,
)
from workflow_dashboard.config import DashboardConfig
from workflow_dashboard.dashboard_html import (
    _bar_chart,
    _closed_by_month_trend,
    _daily_closures_chart,
    _line_chart,
    _table_html,
)
from workflow_dashboard.excel_export import export_workbook
from workflow_dashboard.history import (
    METRIC_DEFINITION_VERSION,
    build_metric_definition,
    load_history,
    store_snapshot,
)
from workflow_dashboard.normalization import normalize_records
from workflow_dashboard.number_format import format_one_decimal, round_one_decimal

GENERATED_AT = pd.Timestamp("2026-07-22T16:00:00Z")


def _half_tie_cases() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "record_id": "case-20",
                "item_label": "Twenty days",
                "workflow_name": "Workflow A",
                "step_name": "Review",
                "status": "In Progress",
                "owner": "Staff Member A",
                "submitted_at": GENERATED_AT - pd.Timedelta(days=20),
                "last_activity_at": GENERATED_AT - pd.Timedelta(days=1),
                "is_terminal": False,
                "age_days": 20.0,
                "idle_days": 1.0,
            },
            {
                "record_id": "case-26-1",
                "item_label": "Twenty-six point one days",
                "workflow_name": "Workflow A",
                "step_name": "Review",
                "status": "In Progress",
                "owner": "Staff Member B",
                "submitted_at": GENERATED_AT - pd.Timedelta(days=26.1),
                "last_activity_at": GENERATED_AT - pd.Timedelta(days=2),
                "is_terminal": False,
                "age_days": 26.1,
                "idle_days": 2.0,
            },
        ]
    )


def _half_tie_analytics():
    return build_analytics(
        _half_tie_cases(),
        config=DashboardConfig(),
        generated_at=GENERATED_AT,
    )


@pytest.mark.parametrize(
    ("value", "rounded", "formatted"),
    [
        (23.05, 23.1, "23.1"),
        (np.float64(23.05), 23.1, "23.1"),
        (23.04, 23.0, "23.0"),
        (-23.05, -23.1, "-23.1"),
        (None, None, "-"),
        (pd.NA, None, "-"),
    ],
)
def test_one_decimal_policy_is_decimal_half_up(value, rounded, formatted) -> None:
    assert round_one_decimal(value) == rounded
    assert format_one_decimal(value) == formatted


def test_half_tie_median_stays_raw_and_displays_consistently() -> None:
    analytics = _half_tie_analytics()
    typical_wait = next(kpi for kpi in analytics.kpis if kpi.key == "typical_wait_time")
    aging = analytics.chart_tables["aging_by_step"]

    assert analytics.snapshot_metrics["median_age_days"] == pytest.approx(23.05)
    assert aging.iloc[0]["age_days"] == pytest.approx(23.05)
    assert typical_wait.value == "23.1d"

    bar = _bar_chart(
        aging,
        "step_name",
        "age_days",
        "Median open-case age by step",
        "Days",
        one_decimal=True,
    )
    assert bar.data[0].x[0] == pytest.approx(23.05)
    assert bar.data[0].text[0] == "23.1"
    assert bar.data[0].customdata[0] == "23.1"

    history = pd.DataFrame({"generated_at": [GENERATED_AT], "median_age_days": [23.05]})
    trend = _line_chart(
        history,
        "median_age_days",
        "Median Open-Case Age Trend",
        "Days",
        one_decimal=True,
    )
    assert trend.data[0].y[0] == pytest.approx(23.05)
    assert trend.data[0].customdata[0][1] == "23.1"


def test_high_cardinality_bar_uses_dynamic_height() -> None:
    table = pd.DataFrame(
        {
            "step_name": [f"Step {index:02}" for index in range(15)],
            "count": [1] * 15,
        }
    )

    bar = _bar_chart(
        table,
        "step_name",
        "count",
        "Open cases by workflow step",
        "Records",
        dynamic_height=True,
    )

    assert len(bar.data[0].y) == 15
    assert bar.layout.height == 520


def test_half_tie_median_round_trips_through_history(tmp_path) -> None:
    analytics = _half_tie_analytics()
    config = DashboardConfig(fixture_path=tmp_path / "fixture.json")
    definition = build_metric_definition(
        config,
        _half_tie_cases(),
        source="fixture",
    )
    db_path = tmp_path / "history.sqlite"

    run_id = store_snapshot(db_path, analytics, definition, source="fixture")
    history, _, _ = load_history(
        db_path,
        source="fixture",
        definition_hash=definition.definition_hash,
    )

    assert definition.metric_definition_version == METRIC_DEFINITION_VERSION == 9
    assert history.loc[history["run_id"] == run_id, "median_age_days"].iloc[
        0
    ] == pytest.approx(23.05)
    with sqlite3.connect(db_path) as connection:
        stored = connection.execute(
            "SELECT median_age_days FROM snapshot_metrics WHERE run_id = ?",
            (run_id,),
        ).fetchone()[0]
    assert stored == pytest.approx(23.05)


def test_normalization_and_csv_retain_raw_age_for_thresholds(tmp_path) -> None:
    age_days = 13.96
    submitted_at = GENERATED_AT - pd.Timedelta(days=age_days)
    records = [
        {
            "packageID": "near-threshold",
            "packageName": "Near threshold",
            "workflowName": "Workflow A",
            "stepName": "Review",
            "status": "In Progress",
            "submittedBy": "Staff Member A",
            "submissionDate": submitted_at.isoformat(),
            "lastActivityDate": (GENERATED_AT - pd.Timedelta(days=1)).isoformat(),
        }
    ]
    config = DashboardConfig(old_age_days=14)

    normalized = normalize_records(records, config, generated_at=GENERATED_AT)
    analytics = build_analytics(
        normalized,
        config=config,
        generated_at=GENERATED_AT,
    )
    csv_path = tmp_path / "normalized.csv"
    _csv_ready(normalized).to_csv(csv_path, index=False)
    csv_frame = pd.read_csv(csv_path)

    assert normalized.iloc[0]["age_days"] == pytest.approx(age_days)
    assert csv_frame.iloc[0]["age_days"] == pytest.approx(age_days)
    assert analytics.exceptions.empty


def test_other_decimal_charts_keep_raw_values_and_use_half_up_labels() -> None:
    closures = pd.DataFrame(
        {
            "closed_date": [pd.Timestamp("2026-07-22")],
            "closed_count": [2],
            "rolling_7_day_avg": [2.25],
        }
    )
    closure_chart = _daily_closures_chart(closures)
    assert closure_chart.data[1].y[0] == pytest.approx(2.25)
    assert closure_chart.data[1].customdata[0] == "2.3"

    monthly = pd.DataFrame(
        {
            "closed_month": ["2026-07"],
            "closed_month_label": ["Jul 2026 (MTD)"],
            "closed_count": [2],
            "median_completion_days": [23.05],
            "completion_valid_count": [2],
            "completion_total_count": [2],
        }
    )
    monthly_chart = _closed_by_month_trend(monthly)
    assert monthly_chart.data[0].customdata[0][0] == "23.1 days"
    assert monthly_chart.data[0].customdata[0][1] == "2/2 valid"

    exception_html = _table_html(
        pd.DataFrame(
            [{"exception_reasons": "Old", "age_days": 23.05, "idle_days": 2.25}]
        )
    )
    assert ">23.1<" in exception_html
    assert ">2.3<" in exception_html
    assert "23.05" not in exception_html


def test_rolling_and_completion_calculations_retain_precision() -> None:
    closures = pd.DataFrame(
        {
            "last_activity_at": pd.to_datetime(
                ["2026-07-15T16:00:00Z", "2026-07-21T16:00:00Z"], utc=True
            ),
            "submitted_at": pd.to_datetime(
                ["2026-06-25T16:00:00Z", "2026-06-25T13:36:00Z"], utc=True
            ),
        }
    )

    daily = _daily_closures(closures, GENERATED_AT)
    monthly = _closed_by_month(closures, GENERATED_AT)

    assert daily.iloc[-1]["rolling_7_day_avg"] == pytest.approx(2 / 7)
    assert monthly.iloc[0]["median_completion_days"] == pytest.approx(23.05)


def test_workbook_rounds_only_exported_values(tmp_path) -> None:
    analytics = _half_tie_analytics()
    history = pd.DataFrame(
        [
            {
                "run_id": "run-1",
                "generated_at": GENERATED_AT,
                "snapshot_date": GENERATED_AT.date(),
                "in_the_works_count": 2,
                "backlog_count": 2,
                "reached_end_today_count": 0,
                "median_age_days": 23.05,
                "exception_count": 2,
            }
        ]
    )
    owners = analytics.chart_tables["owner_workload"].rename(
        columns={"count": "in_the_works_count"}
    )
    owners["run_id"] = "run-1"
    owners["generated_at"] = GENERATED_AT
    owners["snapshot_date"] = GENERATED_AT.date()
    owners["source"] = "fixture"
    workbook_path = tmp_path / "rounding.xlsx"

    export_workbook(
        df=_half_tie_cases(),
        analytics=analytics,
        history=history,
        owner_history=owners,
        output_path=workbook_path,
    )
    workbook = load_workbook(workbook_path, data_only=True)
    history_sheet = workbook["History Trends"]
    history_headers = {
        cell.value: cell.column for cell in next(history_sheet.iter_rows(max_row=1))
    }

    assert workbook["Summary"]["B8"].value == "23.1d"
    assert workbook["Summary"]["E14"].value == pytest.approx(23.1)
    assert history_sheet.cell(
        row=2, column=history_headers["median_age_days"]
    ).value == pytest.approx(23.1)
    assert analytics.snapshot_metrics["median_age_days"] == pytest.approx(23.05)
