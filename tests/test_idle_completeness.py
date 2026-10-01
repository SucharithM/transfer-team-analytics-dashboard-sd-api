import sqlite3

from openpyxl import load_workbook
import pandas as pd
import pytest

from workflow_dashboard.analytics import build_analytics
from workflow_dashboard.config import DashboardConfig, load_config
from workflow_dashboard.dashboard_html import render_dashboard
from workflow_dashboard.excel_export import export_workbook
from workflow_dashboard.history import (
    METRIC_DEFINITION_VERSION,
    HistorySchemaError,
    build_metric_definition,
    connect_history,
    load_history,
    store_snapshot,
)
from workflow_dashboard.normalization import normalize_records


GENERATED_AT = pd.Timestamp("2026-07-22T16:00:00Z")
def _case(index: int, *, idle_days: float | None = 2.0) -> dict:
    last_activity_at = (
        GENERATED_AT - pd.Timedelta(days=idle_days)
        if idle_days is not None
        else pd.NaT
    )
    return {
        "record_id": f"case-{index}",
        "item_label": f"Case {index}",
        "workflow_name": "Workflow A",
        "step_name": "Review",
        "status": "In Progress",
        "owner": "Staff Member A",
        "submitted_at": GENERATED_AT - pd.Timedelta(days=4),
        "last_activity_at": last_activity_at,
        "is_terminal": False,
        "age_days": 4.0,
        "idle_days": idle_days if idle_days is not None else pd.NA,
    }


def _analytics(cases: list[dict], config: DashboardConfig | None = None):
    return build_analytics(
        pd.DataFrame(cases),
        config=config or DashboardConfig(),
        generated_at=GENERATED_AT,
    )


def _idle_kpi(analytics):
    return next(
        kpi for kpi in analytics.kpis if kpi.key == "typical_idle_time"
    )


def test_missing_activity_suppresses_subset_median_and_flags_record() -> None:
    analytics = _analytics([_case(1), _case(2, idle_days=None)])

    assert _idle_kpi(analytics).value == "—"
    assert _idle_kpi(analytics).warning_caption == "1/2 valid · below 95%"
    assert analytics.snapshot_metrics["median_idle_days"] is None
    assert analytics.snapshot_metrics["idle_valid_count"] == 1
    assert analytics.snapshot_metrics["idle_total_count"] == 2
    missing = analytics.exceptions.loc[
        analytics.exceptions["record_id"] == "case-2"
    ].iloc[0]
    assert missing["exception_reasons"] == "Missing last activity"
    counts = analytics.exception_counts.set_index("reason")["count"]
    assert counts["Missing last activity"] == 1


def test_malformed_activity_is_normalized_and_flagged() -> None:
    records = [
        {
            "packageID": "malformed",
            "workflowName": "Workflow A",
            "stepName": "Review",
            "status": "In Progress",
            "submittedBy": "Staff Member A",
            "submissionDate": "2026-07-20T12:00:00Z",
            "lastActivityDate": "not-a-timestamp",
        }
    ]
    normalized = normalize_records(
        records, DashboardConfig(), generated_at=GENERATED_AT
    )
    analytics = build_analytics(
        normalized,
        config=DashboardConfig(),
        generated_at=GENERATED_AT,
    )

    assert pd.isna(normalized.iloc[0]["last_activity_at"])
    assert pd.isna(normalized.iloc[0]["idle_days"])
    assert analytics.exceptions.iloc[0]["exception_reasons"] == (
        "Missing last activity"
    )


@pytest.mark.parametrize(
    ("valid_count", "expected_value", "expected_warning", "expected_median"),
    [
        (20, "2.0d", None, 2.0),
        (19, "2.0d", None, 2.0),
        (18, "—", "18/20 valid · below 95%", None),
        (0, "—", "0/20 valid · below 95%", None),
    ],
)
def test_idle_completeness_threshold(
    valid_count: int,
    expected_value: str,
    expected_warning: str | None,
    expected_median: float | None,
) -> None:
    cases = [
        _case(index, idle_days=2.0 if index < valid_count else None)
        for index in range(20)
    ]
    analytics = _analytics(cases)

    assert _idle_kpi(analytics).value == expected_value
    assert _idle_kpi(analytics).warning_caption == expected_warning
    assert analytics.snapshot_metrics["median_idle_days"] == expected_median


def test_no_open_cases_reports_zero_coverage_without_terminal_exception() -> None:
    terminal = _case(1, idle_days=None)
    terminal.update(step_name="End", status="Completed", is_terminal=True)
    analytics = _analytics([terminal])

    assert _idle_kpi(analytics).value == "—"
    assert _idle_kpi(analytics).warning_caption == "0/0 valid"
    assert analytics.snapshot_metrics["idle_valid_count"] == 0
    assert analytics.snapshot_metrics["idle_total_count"] == 0
    assert analytics.exceptions.empty


def test_missing_activity_combines_with_other_open_case_reasons() -> None:
    case = _case(1, idle_days=None)
    case.update(owner=None, submitted_at=pd.NaT, age_days=pd.NA)
    analytics = _analytics([case])

    assert analytics.exceptions.iloc[0]["exception_reasons"] == (
        "Missing last activity; Missing case submission date; Missing submitter"
    )


def test_coverage_appears_in_dashboard_and_workbook(tmp_path) -> None:
    analytics = _analytics([_case(1), _case(2, idle_days=None)])
    html_path = tmp_path / "dashboard.html"
    workbook_path = tmp_path / "audit.xlsx"
    history = pd.DataFrame()
    owner_history = pd.DataFrame()

    render_dashboard(
        df=analytics.in_the_works,
        analytics=analytics,
        history=history,
        owner_history=owner_history,
        output_path=html_path,
    )
    export_workbook(
        df=analytics.in_the_works,
        analytics=analytics,
        history=history,
        owner_history=owner_history,
        output_path=workbook_path,
    )

    html_text = html_path.read_text(encoding="utf-8")
    assert "1/2 valid · below 95%" in html_text
    assert "Missing last activity" in html_text
    workbook = load_workbook(workbook_path, data_only=True)
    summary_values = {
        workbook["Summary"].cell(row=row, column=1).value:
        workbook["Summary"].cell(row=row, column=2).value
        for row in range(4, 10)
    }
    assert summary_values["Typical Time Since Last Activity"] == "—"
    exception_values = [
        cell.value for row in workbook["Exceptions"].iter_rows() for cell in row
    ]
    assert "Missing last activity" in exception_values


def test_idle_completeness_environment_configuration(monkeypatch, tmp_path) -> None:
    monkeypatch.setenv("WORKFLOW_MIN_WAIT_COMPLETENESS", "0.75")
    monkeypatch.setenv("WORKFLOW_MIN_IDLE_COMPLETENESS", "0.8")
    monkeypatch.setenv("WORKFLOW_MIN_COMPLETION_COMPLETENESS", "0.85")
    config = load_config(env_file=tmp_path / "missing.env")
    assert config.minimum_wait_completeness == pytest.approx(0.75)
    assert config.minimum_idle_completeness == pytest.approx(0.8)
    assert config.minimum_completion_completeness == pytest.approx(0.85)

    monkeypatch.setenv("WORKFLOW_MIN_IDLE_COMPLETENESS", "1.1")
    with pytest.raises(ValueError, match="decimal from 0 through 1"):
        load_config(env_file=tmp_path / "missing.env")

    monkeypatch.setenv("WORKFLOW_MIN_IDLE_COMPLETENESS", "invalid")
    with pytest.raises(ValueError, match="decimal from 0 through 1"):
        load_config(env_file=tmp_path / "missing.env")


def test_threshold_changes_metric_definition() -> None:
    records = pd.DataFrame({"workflow_name": ["Workflow A"]})
    baseline = build_metric_definition(
        DashboardConfig(minimum_idle_completeness=0.95),
        records,
        source="api",
    )
    changed = build_metric_definition(
        DashboardConfig(minimum_idle_completeness=1.0),
        records,
        source="api",
    )

    assert baseline.definition_hash != changed.definition_hash
    assert baseline.metric_definition_version == METRIC_DEFINITION_VERSION == 9


def test_idle_coverage_round_trips_through_history(tmp_path) -> None:
    config = DashboardConfig(fixture_path=tmp_path / "fixture.json")
    cases = pd.DataFrame([_case(1), _case(2, idle_days=None)])
    analytics = build_analytics(
        cases,
        config=config,
        generated_at=GENERATED_AT,
    )
    definition = build_metric_definition(config, cases, source="fixture")
    db_path = tmp_path / "history.sqlite"

    run_id = store_snapshot(db_path, analytics, definition, source="fixture")
    history, _, _ = load_history(
        db_path,
        source="fixture",
        definition_hash=definition.definition_hash,
    )
    stored = history.loc[history["run_id"] == run_id].iloc[0]

    assert stored["idle_valid_count"] == 1
    assert stored["idle_total_count"] == 2
    assert pd.isna(stored["median_idle_days"])


def test_older_versioned_schema_is_rejected_without_mutation(tmp_path) -> None:
    db_path = tmp_path / "history.sqlite"
    with sqlite3.connect(db_path) as connection:
        connection.executescript(
            """
            CREATE TABLE metric_definitions (definition_hash TEXT);
            CREATE TABLE snapshot_metrics (run_id TEXT PRIMARY KEY);
            INSERT INTO snapshot_metrics (run_id) VALUES ('legacy-run');
            CREATE TABLE owner_workload_history (run_id TEXT);
            PRAGMA user_version = 1;
            """
        )

    with pytest.raises(HistorySchemaError, match="Archive or remove"):
        connect_history(db_path)

    with sqlite3.connect(db_path) as connection:
        assert connection.execute("PRAGMA user_version").fetchone()[0] == 1
        row = connection.execute("SELECT run_id FROM snapshot_metrics").fetchone()
    assert row == ("legacy-run",)
