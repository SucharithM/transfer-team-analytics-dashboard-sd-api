from openpyxl import load_workbook
import pandas as pd

from workflow_dashboard.dashboard_html import (
    _latest_daily_history,
    _line_chart,
    _table_html,
    render_dashboard,
)
from workflow_dashboard.excel_export import export_workbook

from .test_analytics_cohorts import GENERATED_AT, build_sample_analytics, sample_cases


def _history_frames(analytics):
    history = pd.DataFrame(
        [
            {
                "run_id": "run-1",
                "generated_at": GENERATED_AT,
                "snapshot_date": GENERATED_AT.date(),
                "backlog_count": analytics.snapshot_metrics["backlog_count"],
                "in_the_works_count": analytics.snapshot_metrics["in_the_works_count"],
                "median_age_days": analytics.snapshot_metrics["median_age_days"],
                "reached_end_today_count": analytics.snapshot_metrics[
                    "reached_end_today_count"
                ],
                "exception_count": analytics.snapshot_metrics["exception_count"],
            }
        ]
    )
    owners = analytics.chart_tables["owner_workload"].rename(
        columns={"count": "in_the_works_count"}
    )
    owners["run_id"] = "run-1"
    owners["generated_at"] = GENERATED_AT
    owners["snapshot_date"] = GENERATED_AT.date()
    owners["source"] = "fixture"
    return history, owners


def test_saved_run_trends_use_latest_eastern_snapshot_per_day() -> None:
    history = pd.DataFrame(
        [
            {
                "run_id": "run-a",
                "generated_at": "2026-07-20T14:00:00Z",
                "median_age_days": 1.0,
                "age_valid_count": 1,
                "age_total_count": 4,
            },
            {
                "run_id": "run-b",
                "generated_at": "2026-07-20T18:00:00Z",
                "median_age_days": 2.0,
                "age_valid_count": 2,
                "age_total_count": 4,
            },
            {
                "run_id": "run-c",
                "generated_at": "2026-07-21T03:30:00Z",
                "median_age_days": 3.0,
                "age_valid_count": 3,
                "age_total_count": 4,
            },
            {
                "run_id": "run-e",
                "generated_at": "2026-07-21T16:00:00Z",
                "median_age_days": 5.0,
                "age_valid_count": 5,
                "age_total_count": 6,
            },
            {
                "run_id": "run-f",
                "generated_at": "2026-07-21T16:00:00Z",
                "median_age_days": pd.NA,
                "age_valid_count": 0,
                "age_total_count": 6,
            },
        ]
    )

    latest = _latest_daily_history(history)
    chart = _line_chart(
        history,
        "median_age_days",
        "Median Open-Case Age Trend",
        "Days",
        one_decimal=True,
        coverage_columns=("age_valid_count", "age_total_count"),
    )

    assert latest["run_id"].tolist() == ["run-c", "run-f"]
    assert list(chart.data[0].x) == [pd.Timestamp("2026-07-20")]
    assert list(chart.data[0].y) == [3.0]
    assert chart.data[0].customdata[0][0] == "2026-07-20 11:30 PM EDT"
    assert chart.data[0].customdata[0][2] == "3/4 valid"
    assert chart.layout.xaxis.title.text == "Date"
    assert list(chart.layout.xaxis.tickvals) == [pd.Timestamp("2026-07-20")]
    assert list(chart.layout.xaxis.range) == [
        pd.Timestamp("2026-07-19"),
        pd.Timestamp("2026-07-21"),
    ]
    assert chart.layout.xaxis.tickformat == "%b %d, %Y"


def test_dashboard_uses_executive_title_and_updated_chart_layout(tmp_path) -> None:
    analytics = build_sample_analytics()
    history, owners = _history_frames(analytics)
    html_path = tmp_path / "dashboard.html"

    render_dashboard(
        df=sample_cases(),
        analytics=analytics,
        history=history,
        owner_history=owners,
        output_path=html_path,
    )
    html_text = html_path.read_text(encoding="utf-8")

    assert analytics.dashboard_title in html_text
    assert "(Prototype)" in html_text
    assert "Run Time" not in html_text
    assert "Run time" not in html_text
    assert "cases have been open for more than 7 days since submission" in html_text
    assert "Daily completions through yesterday" in html_text
    assert "once a full week is available" in html_text
    assert "The current month shows progress so far" in html_text
    assert "Where open cases are in the evaluation process" in html_text
    assert "reached End appear as Completed" in html_text
    assert "Open Cases by Workflow Step" in html_text
    assert "Cases by Status" in html_text
    assert "Open Cases by Submitter" in html_text
    assert "Open Cases by Workflow Step and Case Age" in html_text
    assert "Open Cases: Time and Activity - Risk Overview" in html_text
    assert 'id="plot-age-idle-action-matrix"' in html_text
    assert 'id="age-idle-action-drilldown"' in html_text
    assert 'data-action-bucket="backlog_stale"' in html_text
    assert "Up to 7 Days · Active" in html_text
    assert "Up to 7 Days · Idle" in html_text
    assert 'plot.on("plotly_click"' in html_text
    assert "excluded for missing or invalid timing" in html_text
    assert "Cases without usable dates are included in step totals" in html_text
    assert "Cases Completed — Last 7 Days" in html_text
    assert "Open cases by submitter over time" not in html_text
    assert "Daily Completed Cases" in html_text
    assert "grouped as Other" in html_text
    assert html_text.count('"hole":0.58') == 3
    assert '<div class="current-open-grid">' in html_text
    assert 'id="evaluation-type-summary-table"' in html_text
    assert "Case Overview" in html_text
    assert "Evaluation Cases by Evaluation Type" in html_text
    assert "Evaluation Progress" in html_text
    assert "Total Cases" in html_text
    assert "Package mix" not in html_text
    assert "Total Packages" not in html_text
    assert "Median Open Age" in html_text
    assert "Median Completion Time" in html_text
    assert "Each median includes its valid/total sample" not in html_text
    assert "evaluation-type-table" in html_text
    assert "evaluation-type-section" not in html_text
    assert ".evaluation-type-table-container .table-card-header" not in html_text
    assert ".evaluation-type-table th," in html_text
    assert "min-width: 1080px" in html_text
    assert "Service Performance" not in html_text
    assert "evaluation-type-grid" not in html_text
    assert 'id="plot-evaluation-type-volume"' not in html_text
    assert 'id="plot-evaluation-type-backlog"' not in html_text
    assert 'id="plot-evaluation-type-completion"' not in html_text
    assert "P75" not in html_text
    assert "P90" not in html_text
    assert 'class="chart-container chart-container-wide"' in html_text
    assert ".chart-container-wide .js-plotly-plot" in html_text
    assert "min-width: 720px" in html_text
    assert ".chart-container-wide .card-heading" in html_text
    assert "position: sticky" in html_text
    assert "Steps with older cases appear first" in html_text
    assert "valid/total sample used for its median age" not in html_text
    assert "<th>Package ID</th>" in html_text
    assert "<td>missing-date</td>" in html_text
    assert "<th>Record Id</th>" not in html_text
    assert "grid-template-columns: repeat(7, minmax(0, 1fr))" in html_text
    assert html_text.index("Cases Completed Today") < html_text.index(
        "Total Completed Cases"
    )
    assert html_text.index('aria-labelledby="trends-title"') < html_text.index(
        'aria-labelledby="staff-title"'
    )
    assert html_text.index(
        'aria-labelledby="evaluation-types-title"'
    ) < html_text.index('aria-labelledby="trends-title"')
    assert (
        html_text.index('aria-label="Cases Completed by Month"')
        < html_text.index('aria-label="Daily Completed Cases"')
        < html_text.index('aria-label="Cases Completed — Last 7 Days"')
        < html_text.index('aria-label="Backlog Trend — Open Cases Over 7 Days Old"')
        < html_text.index('aria-label="Median Open-Case Age Trend"')
    )
    assert html_text.index('aria-label="Open Cases by Submitter"') < html_text.index(
        'aria-label="Median Open-Case Age by Workflow Step"'
    )
    assert html_text.index(
        'aria-label="Median Open-Case Age by Workflow Step"'
    ) < html_text.index('aria-label="Open Cases by Workflow Step and Case Age"')
    assert html_text.index(
        'aria-label="Open Cases by Workflow Step and Case Age"'
    ) < html_text.index('aria-label="Open Cases: Time and Activity - Risk Overview"')
    assert 'id="prototype-notice"' not in html_text
    assert "Prototype Notice" not in html_text
    assert (
        "Dashboard metrics are preliminary and may not fully reflect the "
        "underlying data. They will be refined as the prototype evolves. "
        "This report demonstrates the potential of SoftDocs API data "
        "extraction and dashboarding."
    ) not in html_text
    assert ">I understand</button>" not in html_text
    assert "notice.showModal()" not in html_text
    assert "Please note that the dashboard metrics" not in html_text
    assert "Designed and Developed by" not in html_text
    assert "Developed by" in html_text
    assert "<strong>Sucharith Madhusoodana</strong>" in html_text


def test_case_observations_sort_by_age_before_limiting_rows() -> None:
    exceptions = pd.DataFrame(
        [
            {
                "exception_reasons": "Review",
                "record_id": f"case-{age}",
                "age_days": float(age),
            }
            for age in range(102)
        ]
        + [
            {
                "exception_reasons": "Missing case submission date",
                "record_id": "missing-age",
                "age_days": pd.NA,
            }
        ]
    )

    table_html = _table_html(exceptions)

    assert table_html.index("<td>case-101</td>") < table_html.index("<td>case-100</td>")
    assert "<td>case-2</td>" in table_html
    assert "<td>case-1</td>" not in table_html
    assert "<td>case-0</td>" not in table_html
    assert "<td>missing-age</td>" not in table_html


def test_workbook_contains_in_the_works_and_backlog_sheets(tmp_path) -> None:
    analytics = build_sample_analytics()
    history, owners = _history_frames(analytics)
    workbook_path = tmp_path / "audit.xlsx"

    export_workbook(
        df=sample_cases(),
        analytics=analytics,
        history=history,
        owner_history=owners,
        output_path=workbook_path,
    )
    workbook = load_workbook(workbook_path, read_only=True)

    assert workbook["Summary"]["A1"].value == analytics.dashboard_title
    assert workbook["Summary"]["A7"].value == "Total Completed Cases"
    assert workbook["Summary"]["B7"].value == "1"
    assert "In The Works" in workbook.sheetnames
    assert "Backlog" in workbook.sheetnames
    assert "Owner History" in workbook.sheetnames
    assert "Current Backlog" not in workbook.sheetnames
    assert workbook["Owner History"].max_row > 1


def test_workbook_leaves_partial_rolling_windows_blank(tmp_path) -> None:
    analytics = build_sample_analytics()
    history, owners = _history_frames(analytics)
    analytics.chart_tables["daily_closures"] = pd.DataFrame(
        {
            "closed_date": pd.date_range("2026-07-14", periods=7, freq="D"),
            "closed_count": [1] * 7,
            "rolling_7_day_avg": [float("nan")] * 6 + [1.0],
        }
    )
    workbook_path = tmp_path / "rolling-average.xlsx"

    export_workbook(
        df=sample_cases(),
        analytics=analytics,
        history=history,
        owner_history=owners,
        output_path=workbook_path,
    )
    workbook = load_workbook(workbook_path, data_only=True)
    sheet = workbook["Daily Closures"]
    headers = {cell.value: cell.column for cell in next(sheet.iter_rows(max_row=1))}
    average_column = headers["rolling_7_day_avg"]

    assert [
        sheet.cell(row=row, column=average_column).value for row in range(2, 8)
    ] == [None] * 6
    assert sheet.cell(row=8, column=average_column).value == 1.0
    assert sheet.cell(row=8, column=average_column).number_format == "0.0"


def test_workbook_exports_internal_zero_month_without_trailing_report_month(
    tmp_path,
) -> None:
    analytics = build_sample_analytics()
    history, owners = _history_frames(analytics)
    analytics.chart_tables["closed_by_month"] = pd.DataFrame(
        {
            "closed_month": ["2026-01", "2026-02", "2026-03"],
            "closed_month_label": ["Jan 2026", "Feb 2026", "Mar 2026"],
            "closed_count": [1, 0, 1],
            "median_completion_days": [10.0, pd.NA, 12.0],
            "completion_valid_count": [1, 0, 1],
            "completion_total_count": [1, 0, 1],
        }
    )
    workbook_path = tmp_path / "monthly-closures.xlsx"

    export_workbook(
        df=sample_cases(),
        analytics=analytics,
        history=history,
        owner_history=owners,
        output_path=workbook_path,
    )
    workbook = load_workbook(workbook_path, data_only=True)
    sheet = workbook["Closed By Month"]
    rows = list(sheet.iter_rows(values_only=True))
    headers = {value: index for index, value in enumerate(rows[0])}

    assert [row[headers["closed_month"]] for row in rows[1:]] == [
        "2026-01",
        "2026-02",
        "2026-03",
    ]
    assert rows[2][headers["closed_count"]] == 0
    assert rows[2][headers["median_completion_days"]] is None


def test_workbook_exports_monthly_completion_history(tmp_path) -> None:
    analytics = build_sample_analytics()
    history, owners = _history_frames(analytics)
    monthly_history = pd.DataFrame(
        [
            {
                "run_id": "run-1",
                "generated_at": GENERATED_AT,
                "snapshot_date": GENERATED_AT.date(),
                "source": "fixture",
                "definition_hash": "definition",
                "metric_definition_version": 8,
                "app_version": "0.2.0",
                "closed_month": "2026-07",
                "closed_count": 2,
                "median_completion_days": pd.NA,
                "completion_valid_count": 1,
                "completion_total_count": 2,
            }
        ]
    )
    workbook_path = tmp_path / "completion-history.xlsx"

    export_workbook(
        df=sample_cases(),
        analytics=analytics,
        history=history,
        owner_history=owners,
        monthly_completion_history=monthly_history,
        output_path=workbook_path,
    )

    workbook = load_workbook(workbook_path, data_only=True)
    rows = list(workbook["Completion History"].iter_rows(values_only=True))
    headers = {value: index for index, value in enumerate(rows[0])}
    assert rows[1][headers["closed_count"]] == 2
    assert rows[1][headers["median_completion_days"]] is None
    assert rows[1][headers["completion_valid_count"]] == 1
    assert rows[1][headers["completion_total_count"]] == 2
