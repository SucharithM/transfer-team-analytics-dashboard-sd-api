"""Presentation changes must respect configured semantics and audit compatibility."""
import html
from dataclasses import replace

import pandas as pd
import pytest
from openpyxl import load_workbook

from workflow_dashboard.analytics import build_analytics
from workflow_dashboard.config import DashboardConfig
from workflow_dashboard.dashboard_content import COMPLETION_DATE_NOTE
from workflow_dashboard.dashboard_html import render_dashboard
from workflow_dashboard.excel_export import export_workbook
from workflow_dashboard.history import build_metric_definition
from .test_analytics_cohorts import GENERATED_AT, sample_cases


def _render(tmp_path, config, df=None):
    df = sample_cases() if df is None else df
    analytics = build_analytics(df, config=config, generated_at=GENERATED_AT)
    path = tmp_path / 'dashboard.html'
    render_dashboard(df=df, analytics=analytics, history=pd.DataFrame(),
                     owner_history=pd.DataFrame(), output_path=path)
    return analytics, path.read_text()


@pytest.mark.parametrize('title', ['Transfer Credit Evaluation Executive Dashboard', 'Review <Cases> & Timing'])
def test_browser_title_is_exact_configured_title_and_escaped(tmp_path, title):
    _, markup = _render(tmp_path, DashboardConfig(dashboard_title=title))
    assert f'<title>{html.escape(title)}</title>' in markup
    assert 'Dashboard Dashboard' not in markup


def test_nondefault_threshold_copy_matches_backlog_boundary(tmp_path):
    config = DashboardConfig(backlog_age_days=8, old_age_days=20, stale_idle_days=5)
    analytics, markup = _render(tmp_path, config)
    assert 'old-outsider' not in set(analytics.backlog.record_id)  # exactly eight days
    assert 'Backlog Trend — Open Cases Over 8 Days Old' in markup
    assert 'Cases at exactly 8 days are not yet backlog' in markup
    assert 'The 20-day marker highlights longer-running cases' in markup
    assert 'idle means no update for 5 days or more' in markup
    assert 'more than 8 days' in markup
    assert '{backlog_age_days' not in markup
    assert '{old_age_days' not in markup
    assert '{stale_idle_days' not in markup
    assert markup.count(f'class="section-copy completion-date-note">{COMPLETION_DATE_NOTE}') == 3


def test_presentation_context_does_not_change_metrics_or_definition(tmp_path):
    config = DashboardConfig()
    df = sample_cases()
    analytics = build_analytics(df, config=config, generated_at=GENERATED_AT)
    metrics_before = dict(analytics.snapshot_metrics)
    tables_before = {key: table.copy(deep=True) for key, table in analytics.chart_tables.items()}
    definition_before = build_metric_definition(config, df, source='fixture').definition_hash
    render_dashboard(df=df, analytics=analytics, history=pd.DataFrame(),
                     owner_history=pd.DataFrame(), output_path=tmp_path / 'dashboard.html')
    assert analytics.snapshot_metrics == metrics_before
    for key, table in tables_before.items():
        pd.testing.assert_frame_equal(analytics.chart_tables[key], table)
    assert build_metric_definition(replace(config, dashboard_title='New label'), df,
                                   source='fixture').definition_hash == definition_before


def test_summary_copy_changes_without_renaming_audit_tabs_or_columns(tmp_path):
    analytics, markup = _render(tmp_path, DashboardConfig())
    path = tmp_path / 'audit.xlsx'
    export_workbook(df=sample_cases(), analytics=analytics, history=pd.DataFrame(),
                    owner_history=pd.DataFrame(), output_path=path)
    workbook = load_workbook(path, read_only=True)
    assert workbook.sheetnames == ['Summary', 'Normalized Records', 'In The Works',
                                  'Backlog', 'Submitted By Activity', 'Daily Closures',
                                  'Closed By Month', 'Exceptions', 'History Trends',
                                  'Owner History', 'Completion History']
    assert workbook['Summary']['A5'].value == 'Open Cases'
    assert workbook['Summary']['A8'].value == 'Median Open Age'
    assert [c.value for c in next(workbook['Daily Closures'].rows)] == [
        'closed_date', 'closed_count', 'rolling_7_day_avg']
    for label in ('Submitter', 'Workflow Step', 'Case Age (Days)',
                  'Days Since Last Activity', 'Case Submitted', 'Package ID'):
        assert f'<th>{label}</th>' in markup
    assert 'Selected Cases' in markup
    assert 'open cases shown' in markup
    assert 'Median Completion Time' in markup
    assert 'from case submission to the last recorded update' in markup
    assert 'Typical Time in Final Step' not in markup
    for misleading in ('current-step', 'at their current step', 'entering End',
                       'Entered Current Step', 'Time in Step'):
        assert misleading.casefold() not in markup.casefold()
    workbook.close()
