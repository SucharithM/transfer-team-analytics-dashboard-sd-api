"""Excel audit workbook generation."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import dashboard_content as content
from .analytics import AnalyticsBundle
from .number_format import round_one_decimal
from .time_utils import dashboard_timestamp


BLUE = "2563EB"
DARK = "172033"
MUTED = "5F6F85"
BORDER = "D9E2EF"
EXCEL_COLUMN_NAMES = {"record_id": "package_id"}


def _excel_ready(df: pd.DataFrame) -> pd.DataFrame:
    """Apply workbook-specific column labels without changing canonical data names."""

    return df.rename(columns=EXCEL_COLUMN_NAMES)


def _one_decimal_column(column: str) -> bool:
    return column.endswith("_days") or column == "rolling_7_day_avg"


def _excel_value(value: Any, *, one_decimal: bool = False) -> Any:
    if pd.isna(value):
        return None
    if one_decimal:
        return round_one_decimal(value)
    if isinstance(value, pd.Timestamp):
        value = value.floor("us")
        if value.tzinfo is not None:
            value = dashboard_timestamp(value).tz_localize(None)
        return value.to_pydatetime()
    if isinstance(value, str) and len(value) > 32000:
        return value[:32000]
    return value


def _write_dataframe(ws, df: pd.DataFrame, *, start_row: int = 1, start_col: int = 1, table_name: str | None = None) -> None:
    columns = list(df.columns)
    for col_index, column in enumerate(columns, start=start_col):
        cell = ws.cell(row=start_row, column=col_index, value=column)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_index, (_, row) in enumerate(df.iterrows(), start=start_row + 1):
        for col_index, column in enumerate(columns, start=start_col):
            cell = ws.cell(
                row=row_index,
                column=col_index,
                value=_excel_value(
                    row[column], one_decimal=_one_decimal_column(column)
                ),
            )
            cell.alignment = Alignment(vertical="top", wrap_text=column in {"item_label", "source_payload", "exception_reasons"})
            if column.endswith("_at"):
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif _one_decimal_column(column):
                cell.number_format = "0.0"

    max_row = max(start_row + len(df), start_row + 1)
    max_col = start_col + len(columns) - 1
    if table_name and columns:
        ref = f"{get_column_letter(start_col)}{start_row}:{get_column_letter(max_col)}{max_row}"
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        ws.add_table(table)


def _style_sheet(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    thin = Side(style="thin", color=BORDER)
    for row in ws.iter_rows():
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.font = Font(name="Segoe UI", size=10, color=DARK, bold=cell.font.bold)
            if cell.row == 1:
                cell.font = Font(name="Segoe UI", size=10, color="FFFFFF", bold=True)
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        max_len = max((len(str(cell.value)) if cell.value is not None else 0) for cell in column_cells)
        width = min(max(max_len + 2, 12), 48)
        if column_cells[0].value == "source_payload":
            width = 64
        ws.column_dimensions[letter].width = width


def _summary_sheet(ws, analytics: AnalyticsBundle, history: pd.DataFrame) -> None:
    ws.sheet_view.showGridLines = False
    ws["A1"] = analytics.dashboard_title
    ws["A1"].font = Font(name="Segoe UI", size=18, bold=True, color=DARK)
    generated_at = dashboard_timestamp(analytics.generated_at)
    ws["A2"] = (
        f"{content.PAGE.generated_prefix} "
        f"{generated_at.strftime('%Y-%m-%d %H:%M %Z')}"
    )
    ws["A2"].font = Font(name="Segoe UI", size=10, color=MUTED)
    ws.merge_cells("A1:D1")
    ws.merge_cells("A2:D2")

    start_row = 4
    for index, kpi in enumerate(analytics.kpis):
        row = start_row + index
        ws.cell(row=row, column=1, value=kpi.label)
        ws.cell(row=row, column=2, value=kpi.value)
        ws.cell(row=row, column=3, value=kpi.help_text)
        ws.cell(row=row, column=1).font = Font(name="Segoe UI", bold=True, color=DARK)
        ws.cell(row=row, column=2).font = Font(name="Segoe UI", bold=True, color=BLUE)
        ws.cell(row=row, column=3).font = Font(name="Segoe UI", color=MUTED)

    ws["A12"] = content.WORKBOOK.recent_trends_heading
    ws["A12"].font = Font(name="Segoe UI", bold=True, color=DARK)
    trend = history.tail(12).copy()
    if not trend.empty:
        summary_cols = [
            "generated_at",
            "in_the_works_count",
            "backlog_count",
            "reached_end_today_count",
            "median_age_days",
            "age_valid_count",
            "age_total_count",
            "exception_count",
        ]
        _write_dataframe(ws, trend[[column for column in summary_cols if column in trend.columns]], start_row=13, table_name="SummaryTrendTable")
    else:
        ws["A13"] = content.WORKBOOK.empty_trends_message
        ws["A13"].font = Font(name="Segoe UI", color=MUTED)

    widths = {"A": 26, "B": 18, "C": 90, "D": 18, "E": 18}
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width
    thin = Side(style="thin", color=BORDER)
    for row in ws.iter_rows(min_row=4, max_row=10, min_col=1, max_col=3):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor="FFFFFF")
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def export_workbook(
    *,
    df: pd.DataFrame,
    analytics: AnalyticsBundle,
    history: pd.DataFrame,
    owner_history: pd.DataFrame,
    output_path: Path,
    monthly_completion_history: pd.DataFrame | None = None,
) -> Path:
    """Write a manager-readable audit workbook."""

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = content.WORKBOOK.sheet_names["summary"]
    _summary_sheet(summary, analytics, history)
    if monthly_completion_history is None:
        monthly_completion_history = pd.DataFrame(
            columns=[
                "run_id",
                "generated_at",
                "snapshot_date",
                "source",
                "definition_hash",
                "metric_definition_version",
                "app_version",
                "closed_month",
                "closed_count",
                "median_completion_days",
                "completion_valid_count",
                "completion_total_count",
            ]
        )

    sheets = [
        (
            content.WORKBOOK.sheet_names["normalized_records"],
            df,
            "NormalizedRecordsTable",
        ),
        (
            content.WORKBOOK.sheet_names["in_the_works"],
            analytics.in_the_works,
            "InTheWorksTable",
        ),
        (
            content.WORKBOOK.sheet_names["backlog"],
            analytics.backlog,
            "BacklogTable",
        ),
        (
            content.WORKBOOK.sheet_names["submitted_by_activity"],
            analytics.chart_tables["submitted_by_workload"],
            "SubmittedByActivityTable",
        ),
        (
            content.WORKBOOK.sheet_names["daily_closures"],
            analytics.chart_tables["daily_closures"],
            "DailyClosuresTable",
        ),
        (
            content.WORKBOOK.sheet_names["closed_by_month"],
            analytics.chart_tables["closed_by_month"],
            "ClosedByMonthTable",
        ),
        (
            content.WORKBOOK.sheet_names["exceptions"],
            analytics.exceptions,
            "ExceptionsTable",
        ),
        (
            content.WORKBOOK.sheet_names["history_trends"],
            history,
            "HistoryTrendsTable",
        ),
        (
            content.WORKBOOK.sheet_names["owner_history"],
            owner_history,
            "OwnerHistoryTable",
        ),
        (
            content.WORKBOOK.sheet_names["completion_history"],
            monthly_completion_history,
            "CompletionHistoryTable",
        ),
    ]
    for sheet_name, table_df, table_name in sheets:
        ws = workbook.create_sheet(sheet_name)
        write_df = _excel_ready(table_df)
        _write_dataframe(ws, write_df, table_name=table_name)
        _style_sheet(ws)

    for ws in workbook.worksheets:
        # This audit workbook contains no intentional formulas. Preserve every
        # source/configuration string as literal text, including the title and headers.
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0

    workbook.save(output_path)
    return output_path
